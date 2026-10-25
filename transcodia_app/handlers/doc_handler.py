"""Document conversion using the project's independently installed libraries."""
import html
import os
from pathlib import Path
import tempfile

from pypdf import PdfReader, PdfWriter


def read_pdf(path):
    reader = PdfReader(path)
    if reader.is_encrypted or not reader.pages:
        raise ValueError("PDF is empty, encrypted, or unreadable")
    return reader


class DocumentConverter:
    """Extract content and publish complete outputs without helper executables."""

    @staticmethod
    def _publish(destination, write, sources=()):
        temporary = None
        try:
            destination = Path(destination).resolve()
            if any(Path(source).resolve() == destination for source in sources):
                return False
            destination.parent.mkdir(parents=True, exist_ok=True)
            with tempfile.NamedTemporaryFile(suffix=destination.suffix, dir=destination.parent, delete=False) as stream:
                temporary = stream.name
            write(temporary)
            if not os.path.getsize(temporary):
                return False
            os.replace(temporary, destination)
            return True
        except Exception:
            return False
        finally:
            if temporary and os.path.exists(temporary):
                os.unlink(temporary)

    def join_pdfs(self, src_pdfs, dst_pdf):
        try:
            if not src_pdfs:
                return False
            writer = PdfWriter()
            for source in src_pdfs:
                reader = read_pdf(source)
                for page in reader.pages:
                    writer.add_page(page)
            def write(path):
                writer.write(path)
                read_pdf(path)
            return self._publish(dst_pdf, write, src_pdfs)
        except Exception:
            return False

    def convert_pdf_to_xlsx(self, src_pdf, dst_xlsx):
        try:
            import pdfplumber
            import openpyxl
            read_pdf(src_pdf)
            workbook = openpyxl.Workbook()
            workbook.remove(workbook.active)
            with pdfplumber.open(src_pdf) as document:
                for number, page in enumerate(document.pages, 1):
                    sheet = workbook.create_sheet(f"Page {number}")
                    tables = page.extract_tables()
                    if tables:
                        for table in tables:
                            for row in table:
                                sheet.append([str(value) if value is not None else "" for value in row])
                            sheet.append([])
                    else:
                        for line in (page.extract_text() or "").splitlines():
                            sheet.append([line.strip()])
                    for row in sheet:
                        for cell in row:
                            if isinstance(cell.value, str):
                                cell.data_type = "s"
            return self._publish(dst_xlsx, workbook.save, [src_pdf])
        except Exception:
            return False

    def convert_pdf_to_docx(self, src_pdf, dst_docx):
        try:
            from docx import Document
            reader = read_pdf(src_pdf)
            document = Document()
            for number, page in enumerate(reader.pages):
                if number:
                    document.add_page_break()
                document.add_paragraph(page.extract_text() or "")
            return self._publish(dst_docx, document.save, [src_pdf])
        except Exception:
            return False

    def convert_docx_to_pdf(self, src_docx, dst_pdf):
        """Render document text and tables; source page layout is not preserved."""
        try:
            from docx import Document
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table
            document = Document(src_docx)
            style = getSampleStyleSheet()["BodyText"]
            story = []
            for paragraph in document.paragraphs:
                story.extend([Paragraph(html.escape(paragraph.text).replace("\n", "<br/>"), style), Spacer(1, 6)])
            for table in document.tables:
                story.append(Table([[Paragraph(html.escape(cell.text), style) for cell in row.cells] for row in table.rows]))
            if not story:
                story.append(Paragraph(" ", style))
            return self._publish(dst_pdf, lambda path: SimpleDocTemplate(path).build(story), [src_docx])
        except Exception:
            return False

    def convert_pdf_to_txt(self, src_pdf, dst_txt):
        try:
            reader = read_pdf(src_pdf)
            text = "\n\n".join(f"--- Page {number} ---\n{page.extract_text() or ''}"
                               for number, page in enumerate(reader.pages, 1))
            return self._publish(dst_txt, lambda path: Path(path).write_text(text, encoding="utf-8"), [src_pdf])
        except Exception:
            return False

    def convert_pdf_to_html(self, src_pdf, dst_html):
        try:
            reader = read_pdf(src_pdf)
            pages = "\n".join(f'<section id="page-{number}"><pre>{html.escape(page.extract_text() or "")}</pre></section>'
                              for number, page in enumerate(reader.pages, 1))
            text = '<!doctype html><html><head><meta charset="utf-8"><title>Document</title></head><body>' + pages + '</body></html>'
            return self._publish(dst_html, lambda path: Path(path).write_text(text, encoding="utf-8"), [src_pdf])
        except Exception:
            return False

    def convert_to_markdown(self, src_file, dst_md):
        if Path(src_file).suffix.lower() == ".pdf":
            return self.convert_pdf_to_txt(src_file, dst_md)
        return False
