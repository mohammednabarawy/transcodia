from pathlib import Path
import zipfile

from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen.canvas import Canvas
import openpyxl
import pytest

from transcodia_app.handlers.doc_handler import DocumentConverter


@pytest.mark.parametrize("method,extension", [("convert_pdf_to_txt", "txt"), ("convert_pdf_to_html", "html"), ("convert_to_markdown", "md")])
def test_document_exports_retain_text_from_every_page(sample_pdf, tmp_path, method, extension):
    target = tmp_path / f"export.{extension}"
    assert getattr(DocumentConverter(), method)(sample_pdf, str(target))
    text = target.read_text(encoding="utf-8")
    assert "Invoice 123" in text and "Total 456" in text


def test_pdf_excel_is_valid_workbook_with_both_pages(sample_pdf, tmp_path):
    target = tmp_path / "export.xlsx"
    assert DocumentConverter().convert_pdf_to_xlsx(sample_pdf, str(target))
    workbook = openpyxl.load_workbook(target)
    assert len(workbook.worksheets) == 2
    values = [str(cell.value) for sheet in workbook for row in sheet for cell in row if cell.value]
    assert any("Invoice 123" in v for v in values)
    assert any("Total 456" in v for v in values)
    workbook.close()


def test_pdf_table_becomes_separate_excel_cells(tmp_path):
    source = tmp_path / "table.pdf"
    document = Canvas(str(source), pagesize=(300, 200))
    for x in (20, 140, 260):
        document.line(x, 90, x, 180)
    for y in (90, 120, 150, 180):
        document.line(20, y, 260, y)
    for x, y, text in [(30, 160, "Item"), (150, 160, "Amount"), (30, 130, "One"), (150, 130, "123"), (30, 100, "Two"), (150, 100, "456")]:
        document.drawString(x, y, text)
    document.save()
    target = tmp_path / "table.xlsx"
    assert DocumentConverter().convert_pdf_to_xlsx(str(source), str(target))
    workbook = openpyxl.load_workbook(target)
    assert list(workbook.active.values)[:3] == [("Item", "Amount"), ("One", "123"), ("Two", "456")]
    workbook.close()


def test_pdf_text_cannot_turn_into_a_spreadsheet_formula(tmp_path):
    source = tmp_path / "literal.pdf"
    document = Canvas(str(source))
    document.drawString(20, 40, "=1+2")
    document.save()
    target = tmp_path / "literal.xlsx"
    assert DocumentConverter().convert_pdf_to_xlsx(str(source), str(target))
    workbook = openpyxl.load_workbook(target)
    assert workbook.active["A1"].value == "=1+2"
    assert workbook.active["A1"].data_type == "s"
    workbook.close()


@pytest.mark.parametrize("method,extension", [("convert_pdf_to_xlsx", "xlsx"), ("convert_pdf_to_txt", "txt"), ("convert_pdf_to_html", "html"), ("convert_pdf_to_docx", "docx"), ("convert_docx_to_pdf", "pdf"), ("convert_to_markdown", "md")])
def test_invalid_document_fails_without_destroying_existing_output(tmp_path, method, extension):
    source = tmp_path / "corrupt.pdf"
    source.write_bytes(b"broken document")
    target = tmp_path / f"keep.{extension}"
    target.write_bytes(b"previous result")
    assert getattr(DocumentConverter(), method)(str(source), str(target)) is False
    assert target.read_bytes() == b"previous result"


def test_failed_text_conversion_does_not_claim_existing_file_success(tmp_path):
    target = tmp_path / "old.txt"
    target.write_text("old contents")
    assert not DocumentConverter().convert_pdf_to_txt(str(tmp_path / "missing.pdf"), str(target))
    assert target.read_text() == "old contents"


def test_encrypted_pdf_join_fails_cleanly(sample_pdf, tmp_path):
    protected = tmp_path / "protected.pdf"
    writer = PdfWriter(clone_from=sample_pdf)
    writer.encrypt(user_password="password", owner_password="owner", algorithm="AES-256")
    writer.write(protected)
    target = tmp_path / "joined.pdf"
    assert not DocumentConverter().join_pdfs([sample_pdf, str(protected)], str(target))
    assert not target.exists()
    assert not list(tmp_path.glob("tmp*.pdf"))


def test_pdf_to_docx_retains_all_text(sample_pdf, tmp_path):
    target = tmp_path / "export.docx"
    assert DocumentConverter().convert_pdf_to_docx(sample_pdf, str(target))
    with zipfile.ZipFile(target) as archive:
        xml = archive.read("word/document.xml")
        assert b"Invoice 123" in xml and b"Total 456" in xml


def test_docx_to_pdf_uses_our_renderer_and_keeps_text_and_table(tmp_path):
    from docx import Document
    source = tmp_path / "input.docx"
    document = Document()
    document.add_paragraph("Transcodia document test")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Table total 123"
    document.save(source)
    target = tmp_path / "output.pdf"
    assert DocumentConverter().convert_docx_to_pdf(str(source), str(target))
    text = "".join(page.extract_text() for page in PdfReader(target).pages)
    assert "Transcodia document test" in text and "Table total 123" in text


def test_publish_failure_preserves_destination_and_removes_partial_file(tmp_path):
    destination = tmp_path / "keep.txt"
    destination.write_text("previous")
    def fail(path):
        Path(path).write_text("incomplete")
        raise OSError("disk failure")
    assert not DocumentConverter._publish(destination, fail)
    assert destination.read_text() == "previous"
    assert list(tmp_path.iterdir()) == [destination]


def test_html_export_escapes_document_content(tmp_path):
    source = tmp_path / "markup.pdf"
    document = Canvas(str(source))
    document.drawString(20, 40, "<script>alert(1)</script>")
    document.save()
    destination = tmp_path / "document.html"
    assert DocumentConverter().convert_pdf_to_html(str(source), str(destination))
    assert "<script>" not in destination.read_text()
    assert "&lt;script&gt;" in destination.read_text()
